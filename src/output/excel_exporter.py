from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


def save_excel_report(results, output_path):
    workbook = Workbook()
    sheet = workbook.active

    sheet.title = "Invoice Report"

    headers = [
        "File Name",
        "Invoice Number",
        "Vendor",
        "Invoice Date",
        "Subtotal",
        "Tax",
        "Total",
        "Status",
        "Exceptions",
    ]

    sheet.append(headers)

    # Header formatting
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Add invoice data
    for result in results:
        invoice = result["invoice"]
        validation = result["validation"]

        exception_types = ", ".join(
            exception["type"]
            for exception in validation["exceptions"]
        )

        sheet.append([
            result["file_name"],
            invoice["invoice_number"],
            invoice["vendor_name"],
            invoice["invoice_date"],
            invoice["subtotal"],
            invoice["tax"],
            invoice["total"],
            validation["status"],
            exception_types,
        ])

    # Format monetary columns
    for row in sheet.iter_rows(min_row=2):
        for column in [5, 6, 7]:
            row[column - 1].number_format = '#,##0.00'

    # Align important columns
    for row in sheet.iter_rows(min_row=2):
        for column in [5, 6, 7]:
            row[column - 1].alignment = Alignment(horizontal="right")

        row[7].alignment = Alignment(horizontal="center")

    # Highlight failed invoices
    for row in sheet.iter_rows(min_row=2):
        status_cell = row[7]

        if status_cell.value == "FAIL":
            status_cell.font = Font(bold=True)

    # Freeze header row
    sheet.freeze_panes = "A2"

    # Enable filters
    sheet.auto_filter.ref = sheet.dimensions

    # Adjust column widths
    for column_cells in sheet.columns:
        column_letter = get_column_letter(column_cells[0].column)

        max_length = 0

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        sheet.column_dimensions[column_letter].width = min(
            max_length + 2,
            40
        )

    workbook.save(output_path)